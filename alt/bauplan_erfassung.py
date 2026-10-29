#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Bauplan Erfassung v2.1
"""

import os
import re
import sys
import json
import tempfile
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog
from datetime import date

import sv_ttk
from tkinterdnd2 import TkinterDnD, DND_FILES
import fitz
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont


# ═══════════════════════════════════════════════════════════
#  KONFIGURATION
# ═══════════════════════════════════════════════════════════

OUTPUT_DIR = r"\\SERVER\Server A-Z\= Sonja =\Baupläne"

if getattr(sys, "frozen", False):
    APP_DIR = os.path.dirname(sys.executable)
else:
    APP_DIR = os.path.dirname(os.path.abspath(__file__))

KUNDEN_FILE       = os.path.join(APP_DIR, "kunden.txt")
BAUVORHABEN_FILE  = os.path.join(APP_DIR, "bauvorhaben.json")
FONT_FILE         = os.path.join(APP_DIR, "PTSansNarrow-Bold.ttf")

FARBE_OPTIONEN  = ["Farbig", "S/W"]
SPALTEN_HEADERS = ["Datum", "Bauvorhaben", "Dateiname", "Farbig/S/W",
                   "Anzahl", "Breite (mm)", "Höhe (mm)", "Gesamtfläche (m²)"]
SPALTEN_BREITEN = [12, 28, 42, 12, 8, 12, 10, 18]

PDF_ORANGE  = colors.HexColor("#F07820")
PDF_GRAY    = colors.HexColor("#777777")
PDF_LGRAY   = colors.HexColor("#CCCCCC")
PDF_TEXT    = colors.HexColor("#1A1A1A")
PDF_ROWTINT = colors.HexColor("#FAFAFA")


# ═══════════════════════════════════════════════════════════
#  DATENHALTUNG
# ═══════════════════════════════════════════════════════════

def load_kunden():
    if not os.path.exists(KUNDEN_FILE):
        messagebox.showwarning(
            "kunden.txt fehlt",
            f"Nicht gefunden:\n{KUNDEN_FILE}\n\n"
            "Format: Ein Kunde pro Zeile."
        )
        return []
    with open(KUNDEN_FILE, encoding="utf-8") as f:
        return sorted(line.strip() for line in f if line.strip())


def save_kunden(kunden):
    with open(KUNDEN_FILE, "w", encoding="utf-8") as f:
        f.write("\n".join(sorted(kunden)))


def load_bv_db():
    if not os.path.exists(BAUVORHABEN_FILE):
        return {}
    try:
        with open(BAUVORHABEN_FILE, encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}


def save_bv_db(db):
    with open(BAUVORHABEN_FILE, "w", encoding="utf-8") as f:
        json.dump(db, f, ensure_ascii=False, indent=2)


def get_bauvorhaben(kunde):
    return load_bv_db().get(kunde, [])


def remember_bauvorhaben(kunde, bv):
    bv = bv.strip()
    if not bv:
        return
    db  = load_bv_db()
    lst = db.get(kunde, [])
    if bv in lst:
        lst.remove(bv)
    lst.insert(0, bv)
    db[kunde] = lst[:20]
    save_bv_db(db)


# ═══════════════════════════════════════════════════════════
#  PDF-GENERIERUNG
# ═══════════════════════════════════════════════════════════

def _register_font():
    """PT Sans Narrow Bold registrieren. Gibt den Font-Namen zurück."""
    if os.path.exists(FONT_FILE):
        try:
            pdfmetrics.registerFont(TTFont("PTSansNarrow-Bold", FONT_FILE))
            return "PTSansNarrow-Bold"
        except Exception:
            pass
    return "Helvetica-Bold"


def _truncate(text, max_chars):
    return text if len(text) <= max_chars else text[:max_chars - 1] + "…"


def create_pdf(kunde, bv, datum, rows):
    """
    rows: list of dicts — filename, farbe, anzahl, breite, hoehe, flaeche
    Erzeugt eine temporäre PDF-Datei und gibt den Pfad zurück.
    """
    brand_font = _register_font()

    fd, path = tempfile.mkstemp(suffix=".pdf", prefix="printandmore_")
    os.close(fd)

    page_w, page_h = A4
    margin   = 18 * mm
    usable_w = page_w - 2 * margin

    c = rl_canvas.Canvas(path, pagesize=A4)
    c.setTitle(f"Bauplan – {kunde}")

    y = page_h - margin

    # ── Markenname ────────────────────────────────────────
    c.setFont(brand_font, 24)
    c.setFillColor(PDF_ORANGE)
    c.drawString(margin, y - 6 * mm, "printandmore")
    y -= 13 * mm

    # ── Trennlinie ────────────────────────────────────────
    c.setStrokeColor(PDF_ORANGE)
    c.setLineWidth(0.75)
    c.line(margin, y, page_w - margin, y)
    y -= 8 * mm

    # ── Info-Block ────────────────────────────────────────
    label_x = margin
    value_x = margin + 32 * mm
    row_h   = 5.5 * mm

    c.setFont("Helvetica", 8.5)
    for label, value in [("Kunde", kunde), ("Bauvorhaben", bv), ("Datum", datum)]:
        c.setFillColor(PDF_GRAY)
        c.drawString(label_x, y, label)
        c.setFillColor(PDF_TEXT)
        c.drawString(value_x, y, value or "–")
        y -= row_h

    y -= 9 * mm

    # ── Tabelle ───────────────────────────────────────────
    # Spaltenbreiten in Punkten
    col_w = [
        usable_w * 0.38,   # Dateiname
        usable_w * 0.11,   # Farbe
        usable_w * 0.08,   # Anzahl
        usable_w * 0.11,   # Breite
        usable_w * 0.11,   # Höhe
        usable_w * 0.13,   # m²
    ]
    col_x = [margin]
    for w in col_w[:-1]:
        col_x.append(col_x[-1] + w)

    headers  = ["Dateiname", "Farbe", "Anz.", "Breite", "Höhe", "m²"]
    row_h_t  = 6 * mm
    head_h   = 6.5 * mm

    # Tabellenüberschrift
    c.setFont("Helvetica-Bold", 8)
    c.setFillColor(PDF_GRAY)
    aligns = ["L", "L", "R", "R", "R", "R"]
    for i, (hdr, x, w, al) in enumerate(zip(headers, col_x, col_w, aligns)):
        if al == "R":
            c.drawRightString(x + w - 1 * mm, y, hdr)
        else:
            c.drawString(x + 1 * mm, y, hdr)
    y -= 4 * mm

    # Linie unter Header
    c.setStrokeColor(PDF_LGRAY)
    c.setLineWidth(0.5)
    c.line(margin, y, page_w - margin, y)
    y -= row_h_t

    # Datenzeilen
    c.setFont("Helvetica", 8.5)
    for idx, r in enumerate(rows):
        # Alternierender Hintergrund
        if idx % 2 == 1:
            c.setFillColor(PDF_ROWTINT)
            c.rect(margin, y - 1.5 * mm, usable_w,
                   row_h_t, fill=1, stroke=0)

        c.setFillColor(PDF_TEXT)
        cells = [
            ("L", _truncate(r["filename"], 38)),
            ("L", r["farbe"]),
            ("R", str(r["anzahl"])),
            ("R", str(r["breite"])),
            ("R", str(r["hoehe"])),
            ("R", f'{r["flaeche"]:.4f}'),
        ]
        for (al, text), x, w in zip(cells, col_x, col_w):
            if al == "R":
                c.drawRightString(x + w - 1 * mm, y, text)
            else:
                c.drawString(x + 1 * mm, y, text)
        y -= row_h_t

    # Abschlusslinie
    c.setStrokeColor(PDF_LGRAY)
    c.setLineWidth(0.5)
    c.line(margin, y + row_h_t - 1.5 * mm, page_w - margin, y + row_h_t - 1.5 * mm)

    # ── Kassenblock ───────────────────────────────────────
    y -= 12 * mm

    m2_farbig  = sum(r["flaeche"] for r in rows if r["farbe"] == "Farbig")
    m2_sw      = sum(r["flaeche"] for r in rows if r["farbe"] == "S/W")
    anz_gesamt = sum(r["anzahl"]  for r in rows)

    block_w = 68 * mm
    block_h = 22 * mm
    block_x = page_w - margin - block_w
    block_y = y - block_h

    # Rahmen
    c.setStrokeColor(PDF_LGRAY)
    c.setLineWidth(0.6)
    c.roundRect(block_x, block_y, block_w, block_h, 2 * mm, fill=0, stroke=1)

    # Kasseninhalt
    lbl_x  = block_x + 4 * mm
    val_x  = block_x + block_w - 4 * mm
    line_y = y - 5 * mm

    kassen_rows = [
        ("m² farbig",    f"{m2_farbig:.4f}"),
        ("m² s/w",       f"{m2_sw:.4f}"),
        ("Anzahl Pläne", str(anz_gesamt)),
    ]
    for label, value in kassen_rows:
        c.setFont("Helvetica", 8.5)
        c.setFillColor(PDF_GRAY)
        c.drawString(lbl_x, line_y, label)
        c.setFont("Helvetica-Bold", 9)
        c.setFillColor(PDF_TEXT)
        c.drawRightString(val_x, line_y, value)
        line_y -= 5.5 * mm

    c.save()
    return path


# ═══════════════════════════════════════════════════════════
#  PDF ÖFFNEN
# ═══════════════════════════════════════════════════════════

def open_pdf(path):
    """PDF im Standard-PDF-Viewer öffnen (Windows)."""
    try:
        os.startfile(path)
    except AttributeError:
        # Fallback für Nicht-Windows (Entwicklung/Test)
        import subprocess
        subprocess.Popen(["xdg-open", path])


# ═══════════════════════════════════════════════════════════
#  PDF-HILFSFUNKTIONEN
# ═══════════════════════════════════════════════════════════

def parse_drop_paths(raw):
    found = re.findall(r'\{([^}]+)\}|(\S+)', raw)
    return [b or p for b, p in found if (b or p).lower().endswith(".pdf")]


def pt_to_mm(pt):
    return round(pt * 25.4 / 72)


def detect_color(doc):
    try:
        pix  = doc[0].get_pixmap(dpi=36, colorspace=fitz.csRGB)
        data = bytearray(pix.samples)
        for i in range(0, len(data) - 2, 30):
            r, g, b = data[i], data[i + 1], data[i + 2]
            if abs(int(r) - int(g)) > 15 or abs(int(g) - int(b)) > 15:
                return "Farbig"
        return "S/W"
    except Exception:
        return "S/W"


def get_pdf_info(filepath):
    doc    = fitz.open(filepath)
    rect   = doc[0].rect
    breite = pt_to_mm(rect.width)
    hoehe  = pt_to_mm(rect.height)
    farbe  = detect_color(doc)
    doc.close()
    return breite, hoehe, farbe


# ═══════════════════════════════════════════════════════════
#  EXCEL
# ═══════════════════════════════════════════════════════════

def save_to_excel(kunde, rows):
    if not os.path.isdir(OUTPUT_DIR):
        raise FileNotFoundError(
            f"Netzlaufwerk nicht erreichbar:\n{OUTPUT_DIR}"
        )
    filepath = os.path.join(OUTPUT_DIR, f"{kunde}.xlsx")
    if os.path.exists(filepath):
        wb = load_workbook(filepath)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Baupläne"
        for col, (h, w) in enumerate(zip(SPALTEN_HEADERS, SPALTEN_BREITEN), 1):
            cell           = ws.cell(row=1, column=col, value=h)
            cell.font      = Font(bold=True, color="FFFFFF")
            cell.fill      = PatternFill("solid", fgColor="2D5FA0")
            cell.alignment = Alignment(horizontal="center")
            ws.column_dimensions[get_column_letter(col)].width = w
        ws.freeze_panes = "A2"
    for row in rows:
        ws.append(row)
    wb.save(filepath)
    return filepath


# ═══════════════════════════════════════════════════════════
#  GUI
# ═══════════════════════════════════════════════════════════

class BauplanApp(TkinterDnD.Tk):

    def __init__(self):
        super().__init__()
        self.title("Bauplan Erfassung")
        self.resizable(True, False)
        self.minsize(620, 420)
        self.pdf_rows = []
        self.kunden   = load_kunden()
        sv_ttk.set_theme("light")
        self._build_ui()

    # ── Aufbau ─────────────────────────────────────────────────────────────────

    def _build_ui(self):
        outer = ttk.Frame(self, padding=18)
        outer.pack(fill="both", expand=True)

        ttk.Label(
            outer, text="Bauplan Erfassung",
            font=("Segoe UI", 16, "bold")
        ).pack(anchor="w", pady=(0, 16))

        # ── Auftrag ──────────────────────────────────────────────────────────
        auftrag = ttk.LabelFrame(outer, text="Auftrag", padding=(12, 10))
        auftrag.pack(fill="x", pady=(0, 14))

        row_k = ttk.Frame(auftrag)
        row_k.pack(fill="x", pady=(0, 8))
        ttk.Label(row_k, text="Kunde", width=14, anchor="w").pack(side="left")
        self.kunde_var = tk.StringVar()
        self.kunde_cb  = ttk.Combobox(
            row_k, textvariable=self.kunde_var,
            values=self.kunden, state="readonly", width=34
        )
        self.kunde_cb.pack(side="left", padx=(0, 8))
        if self.kunden:
            self.kunde_cb.set(self.kunden[0])
        self.kunde_cb.bind("<<ComboboxSelected>>", self._on_kunde_changed)
        ttk.Button(row_k, text="+ Neu", width=8,
                   command=self._add_kunde).pack(side="left")

        row_bv = ttk.Frame(auftrag)
        row_bv.pack(fill="x")
        ttk.Label(row_bv, text="Bauvorhaben", width=14, anchor="w").pack(side="left")
        self.bv_var = tk.StringVar()
        self.bv_cb  = ttk.Combobox(
            row_bv, textvariable=self.bv_var, values=[], width=44
        )
        self.bv_cb.pack(side="left")
        self._on_kunde_changed()

        # ── Drop-Zone ────────────────────────────────────────────────────────
        self.drop_zone = tk.Label(
            outer,
            text="PDF-Dateien hier ablegen  —  oder klicken zum Auswählen",
            font=("Segoe UI", 10),
            fg="#4A6FA5", bg="#EEF3FA",
            relief="flat", bd=0,
            pady=22, cursor="hand2",
        )
        self.drop_zone.pack(fill="x", pady=(0, 4))
        self.drop_zone.config(
            highlightbackground="#A8C0DE",
            highlightcolor="#4A6FA5",
            highlightthickness=2
        )
        self.drop_zone.drop_target_register(DND_FILES)
        self.drop_zone.dnd_bind("<<Drop>>", self._on_drop)
        self.drop_zone.bind("<Button-1>", lambda e: self.select_pdfs())
        self.drop_zone.bind("<Enter>",
            lambda e: self.drop_zone.config(bg="#DCE8F5", fg="#1A3F75"))
        self.drop_zone.bind("<Leave>",
            lambda e: self.drop_zone.config(bg="#EEF3FA", fg="#4A6FA5"))

        self.status_label = ttk.Label(outer, text="", foreground="#666666")
        self.status_label.pack(anchor="w", pady=(4, 10))

        # ── Tabelle ──────────────────────────────────────────────────────────
        self.table_frame = ttk.Frame(outer)
        self.table_frame.pack(fill="both", expand=True)
        self._draw_table_header()

        # ── Kassenwerte ──────────────────────────────────────────────────────
        self.kasse_frame = ttk.LabelFrame(outer, text="Kassenwerte", padding=(12, 8))
        self.kasse_frame.pack(fill="x", pady=(12, 0))

        self.lbl_m2_farbig  = ttk.Label(self.kasse_frame, text="m² farbig  —",
                                         font=("Segoe UI", 10))
        self.lbl_m2_sw      = ttk.Label(self.kasse_frame, text="m² s/w  —",
                                         font=("Segoe UI", 10))
        self.lbl_anz        = ttk.Label(self.kasse_frame, text="Anzahl Pläne  —",
                                         font=("Segoe UI", 10))

        self.lbl_m2_farbig.pack(side="left", padx=(0, 28))
        self.lbl_m2_sw.pack(side="left", padx=(0, 28))
        self.lbl_anz.pack(side="left")

        # ── Buttons ──────────────────────────────────────────────────────────
        row_bottom = ttk.Frame(outer)
        row_bottom.pack(fill="x", pady=(14, 0))

        self.save_btn = ttk.Button(
            row_bottom, text="In Excel speichern",
            command=self.save, state="disabled"
        )
        self.save_btn.pack(side="right", padx=(8, 0))

        self.pdf_btn = ttk.Button(
            row_bottom, text="Als PDF erstellen",
            command=self._create_pdf, state="disabled"
        )
        self.pdf_btn.pack(side="right")

    def _draw_table_header(self):
        for w in self.table_frame.winfo_children():
            w.destroy()
        self.pdf_rows.clear()
        cols = [("Dateiname", 38), ("Farbe", 10), ("Anzahl", 7),
                ("Breite (mm)", 11), ("Höhe (mm)", 11)]
        for c, (h, w) in enumerate(cols):
            ttk.Label(
                self.table_frame, text=h, width=w, anchor="w"
            ).grid(row=0, column=c, padx=(0, 6), pady=(0, 6), sticky="w")
        self._update_summary()

    # ── Ereignisse ─────────────────────────────────────────────────────────────

    def _update_summary(self, *_):
        """Kassenwerte live neu berechnen und anzeigen."""
        if not hasattr(self, "lbl_m2_farbig"):
            return
        m2_farbig = m2_sw = 0.0
        anz = 0
        for r in self.pdf_rows:
            try:
                anzahl = int(r["anzahl_var"].get())
                breite = int(r["breite_var"].get())
                hoehe  = int(r["hoehe_var"].get())
            except (ValueError, tk.TclError):
                continue
            fl = (breite * hoehe * anzahl) / 1_000_000
            if r["farbe_var"].get() == "Farbig":
                m2_farbig += fl
            else:
                m2_sw += fl
            anz += anzahl

        if self.pdf_rows:
            self.lbl_m2_farbig.config(text=f"m² farbig  {m2_farbig:.4f}")
            self.lbl_m2_sw.config(    text=f"m² s/w  {m2_sw:.4f}")
            self.lbl_anz.config(      text=f"Anzahl Pläne  {anz}")
        else:
            self.lbl_m2_farbig.config(text="m² farbig  —")
            self.lbl_m2_sw.config(    text="m² s/w  —")
            self.lbl_anz.config(      text="Anzahl Pläne  —")

    def _on_kunde_changed(self, event=None):
        kunde   = self.kunde_var.get()
        bv_list = get_bauvorhaben(kunde)
        self.bv_cb["values"] = bv_list
        self.bv_var.set(bv_list[0] if bv_list else "")

    def _add_kunde(self):
        name = simpledialog.askstring(
            "Neuer Kunde", "Kundenname:", parent=self
        )
        if not name or not name.strip():
            return
        name = name.strip()
        if name in self.kunden:
            messagebox.showinfo("Hinweis", f'"{name}" ist bereits vorhanden.')
            return
        self.kunden.append(name)
        save_kunden(self.kunden)
        self.kunden = load_kunden()
        self.kunde_cb["values"] = self.kunden
        self.kunde_var.set(name)
        self._on_kunde_changed()

    def _on_drop(self, event):
        files = parse_drop_paths(event.data)
        if not files:
            messagebox.showwarning("Hinweis", "Keine PDF-Dateien erkannt.")
            return
        self._load_files(files)

    def select_pdfs(self):
        files = filedialog.askopenfilenames(
            title="Baupläne auswählen",
            filetypes=[("PDF-Dateien", "*.pdf")]
        )
        if files:
            self._load_files(list(files))

    def _load_files(self, files):
        self._draw_table_header()
        self.status_label.config(text="Wird gelesen …")
        self.update()
        geladen = 0
        for row_idx, filepath in enumerate(files, 1):
            filename = os.path.basename(filepath)
            try:
                breite, hoehe, farbe = get_pdf_info(filepath)
            except Exception as e:
                messagebox.showerror("Lesefehler", f"{filename}:\n{e}")
                continue

            farbe_var  = tk.StringVar(value=farbe)
            anzahl_var = tk.IntVar(value=1)
            breite_var = tk.StringVar(value=str(breite))
            hoehe_var  = tk.StringVar(value=str(hoehe))

            ttk.Label(self.table_frame, text=filename,
                      anchor="w", width=38).grid(
                row=row_idx, column=0, padx=(0, 6), pady=3, sticky="w")
            ttk.Combobox(self.table_frame, textvariable=farbe_var,
                         values=FARBE_OPTIONEN, state="readonly", width=8).grid(
                row=row_idx, column=1, padx=(0, 6), pady=3)
            ttk.Spinbox(self.table_frame, textvariable=anzahl_var,
                        from_=1, to=999, width=5).grid(
                row=row_idx, column=2, padx=(0, 6), pady=3)
            ttk.Entry(self.table_frame, textvariable=breite_var, width=11).grid(
                row=row_idx, column=3, padx=(0, 6), pady=3)
            ttk.Entry(self.table_frame, textvariable=hoehe_var, width=11).grid(
                row=row_idx, column=4, pady=3)

            self.pdf_rows.append({
                "filename":   filename,
                "farbe_var":  farbe_var,
                "anzahl_var": anzahl_var,
                "breite_var": breite_var,
                "hoehe_var":  hoehe_var,
            })
            # Live-Update der Kassenwerte bei jeder Änderung
            for var in (farbe_var, anzahl_var, breite_var, hoehe_var):
                var.trace_add("write", self._update_summary)
            geladen += 1

        if geladen:
            self.status_label.config(text=f"{geladen} Datei(en) geladen")
            self.save_btn.config(state="normal")
            self.pdf_btn.config(state="normal")
            self._update_summary()
        else:
            self.status_label.config(text="Keine Dateien geladen")
            self.save_btn.config(state="disabled")
            self.pdf_btn.config(state="disabled")

    def _collect_rows(self):
        """Liest aktuelle Tabellenwerte und gibt sie als Liste von Dicts zurück."""
        rows = []
        for r in self.pdf_rows:
            try:
                anzahl = int(r["anzahl_var"].get())
                breite = int(r["breite_var"].get())
                hoehe  = int(r["hoehe_var"].get())
            except (ValueError, tk.TclError):
                messagebox.showerror(
                    "Eingabefehler",
                    f"Ungültige Zahl bei:\n{r['filename']}"
                )
                return None
            flaeche = round((breite * hoehe * anzahl) / 1_000_000, 4)
            rows.append({
                "filename": r["filename"],
                "farbe":    r["farbe_var"].get(),
                "anzahl":   anzahl,
                "breite":   breite,
                "hoehe":    hoehe,
                "flaeche":  flaeche,
            })
        return rows

    def _create_pdf(self):
        rows = self._collect_rows()
        if rows is None:
            return
        kunde = self.kunde_var.get()
        bv    = self.bv_var.get().strip()
        heute = date.today().strftime("%d.%m.%Y")
        try:
            path = create_pdf(kunde, bv, heute, rows)
            open_pdf(path)
        except Exception as e:
            messagebox.showerror("PDF-Fehler", str(e))

    def save(self):
        rows = self._collect_rows()
        if rows is None:
            return
        kunde = self.kunde_var.get()
        bv    = self.bv_var.get().strip()
        heute = date.today().strftime("%d.%m.%Y")

        excel_rows = [
            [heute, bv, r["filename"], r["farbe"],
             r["anzahl"], r["breite"], r["hoehe"], r["flaeche"]]
            for r in rows
        ]
        try:
            filepath = save_to_excel(kunde, excel_rows)
            remember_bauvorhaben(kunde, bv)
            messagebox.showinfo(
                "Gespeichert",
                f"{len(rows)} Eintrag/Einträge gespeichert:\n{filepath}"
            )
            self._draw_table_header()
            self.status_label.config(text="")
            self.save_btn.config(state="disabled")
            self.pdf_btn.config(state="disabled")
            self._on_kunde_changed()
        except Exception as e:
            messagebox.showerror("Speicherfehler", str(e))


# ═══════════════════════════════════════════════════════════
#  START
# ═══════════════════════════════════════════════════════════

if __name__ == "__main__":
    app = BauplanApp()
    app.mainloop()
